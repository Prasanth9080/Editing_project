# views.py
import datetime
from django.shortcuts import render, redirect
from django.contrib import messages
from django.contrib.auth import login, logout
from rest_framework_simplejwt.tokens import RefreshToken
from .models import User
import random, string

otp_storage = {}

def generate_otp():
    return ''.join(random.choices(string.digits, k=6))  # 6-digit OTP


from django.shortcuts import render, redirect
from django.contrib import messages
from django.contrib.auth.models import BaseUserManager
from rest_framework_simplejwt.tokens import RefreshToken

from .models import User  # your custom user model

def signup_view(request):
    if request.method == 'POST':
        name = request.POST.get('username')
        phone = request.POST.get('phone_number')
        email = request.POST.get('email')

        # Server-side validation
        if len(phone) != 10 or not phone.isdigit():
            messages.error(request, "Enter a valid 10-digit phone number")
            return redirect('signup')

        if User.objects.filter(phone_number=phone).exists():
            messages.error(request, "Phone number already registered")
            return redirect('signup')

        if User.objects.filter(email=email).exists():
            messages.error(request, "Email already registered")
            return redirect('signup')

        # Generate random password using BaseUserManager
        # random_password = BaseUserManager().make_random_password()

        # Create user
        user = User.objects.create_user(
            username=name,
            email=email,
            phone_number=phone,
            # password=random_password
        )

        # Generate JWT token
        refresh = RefreshToken.for_user(user)
        user.jwt_token = str(refresh.access_token)
        user.save()

        messages.success(request, "Signup successful! Please login.")
        return redirect('login')

    return render(request, 'signup.html')

def login_view(request):
    if request.method == 'POST':
        phone = request.POST.get('phone_number')

        try:
            user = User.objects.get(phone_number=phone)

            if not user.jwt_token:
                # Generate new token if missing
                refresh = RefreshToken.for_user(user)
                user.jwt_token = str(refresh.access_token)
                user.save()

            otp = generate_otp()
            otp_storage[phone] = otp
            request.session['phone_number'] = phone

            print(f"\n🔐 OTP for {phone}: {otp}\n")  # Print clearly to terminal

            messages.success(request, "OTP sent to your number.")
            return redirect('verify_otp')

        except User.DoesNotExist:
            messages.error(request, "This number is not registered. Please sign up.")
            return redirect('signup')

    return render(request, 'login.html')


def verify_otp_view(request):
    phone = request.session.get('phone_number')
    if not phone:
        return redirect('login')

    if request.method == 'POST':
        input_otp = request.POST.get('otp')
        expected_otp = otp_storage.get(phone)

        if input_otp == expected_otp:
            try:
                user = User.objects.get(phone_number=phone)
            except User.DoesNotExist:
                messages.error(request, "User not found.")
                return redirect('login')

            login(request, user)

            # Generate access token
            refresh = RefreshToken.for_user(user)
            access_token = str(refresh.access_token)
            request.session['access_token'] = access_token

            response = redirect('formpage')
            response.set_cookie('auth_token', access_token)
            return response
        else:
            messages.error(request, "Invalid OTP.")

    return render(request, 'verify_otp.html')


def logout_view(request):
    logout(request)
    messages.success(request, "Logged out successfully.")
    response = redirect('login')
    response.delete_cookie('jwt_token')
    return response



################ delete option only acccess in main user
################ otherwise normal user delete the record only delete(hide) the paricular role
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import MyKYC, SubKYC, BondImage
from django.contrib.auth import get_user_model

User = get_user_model()

@login_required
def form_page(request):
    user = request.user
    user_id = request.GET.get('user_id')
    selected_user = None

    hidden_my_kyc_ids = request.session.get('hidden_my_kyc', [])
    hidden_sub_kyc_ids = request.session.get('hidden_sub_kyc', [])

    # Determine which user's data to show
    if user.is_main_user or user.is_sub_mainuser:
        if user_id:
            try:
                selected_user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                messages.error(request, "Selected user does not exist.")
                return redirect("formpage")
        else:
            selected_user = user
    else:
        selected_user = user

    # Get KYC and Sub-KYC for the selected user
    my_kyc_list = MyKYC.objects.filter(created_by=selected_user).exclude(id__in=hidden_my_kyc_ids)

    if user.is_main_user or user.is_sub_mainuser:
        # Show all Sub-KYCs created by or for the selected user
        sub_kyc_list = SubKYC.objects.filter(
            user=selected_user
        ) | SubKYC.objects.filter(
            created_by=selected_user
        )
        sub_kyc_list = sub_kyc_list.exclude(id__in=hidden_sub_kyc_ids).distinct()
    else:
        # Normal user only sees their own
        sub_kyc_list = SubKYC.objects.filter(user=user).exclude(id__in=hidden_sub_kyc_ids)

    users = User.objects.filter(is_superuser=False, is_main_user=False, is_sub_mainuser=False).exclude(id=user.id)

    # Handle form submission
    if request.method == "POST":
        name = request.POST.get('name')
        age = request.POST.get('age')
        fathername = request.POST.get('fathername')
        mobile = request.POST.get('mobile_number')
        aadhar = request.POST.get('aadhar_number')
        aadhar_front_image = request.FILES.get('aadhar_front_image')
        aadhar_back_image = request.FILES.get('aadhar_back_image')
        address = request.POST.get('address')
        profession = request.POST.get('profession')
        contactSH = request.POST.get('contactSH')
        nameSH = request.POST.get('nameSH')
        investmentamt = request.POST.get('investmentamt')
        passportphoto = request.FILES.get('passportphoto')
        bonds = request.FILES.getlist('bonds')

        is_sub_kyc = 'is_sub_kyc' in request.POST
        data_for_user = user

        if is_sub_kyc and (user.is_main_user or user.is_sub_mainuser):
            data_for_user_id = request.POST.get('data_for_user')
            try:
                data_for_user = User.objects.get(id=data_for_user_id)
            except User.DoesNotExist:
                messages.error(request, "Invalid user selected.")
                return redirect("formpage")

        if not all([name, mobile, aadhar, aadhar_front_image, aadhar_back_image, address]):
            messages.error(request, "All required fields must be filled.")
            return redirect("formpage")

        try:
            investmentamt = int(investmentamt) if investmentamt else 0
        except ValueError:
            messages.error(request, "Investment amount must be numeric.")
            return redirect("formpage")

        if is_sub_kyc:
            subkyc = SubKYC.objects.create(
                created_by=user,
                user=data_for_user,
                name=name,
                age=age,
                fathername=fathername,
                mobile_number=mobile,
                aadhar_number=aadhar,
                aadhar_front_image=aadhar_front_image,
                aadhar_back_image=aadhar_back_image,
                address=address,
                profession=profession,
                contactSH=contactSH,
                nameSH=nameSH,
                investmentamt=investmentamt,
                passportphoto=passportphoto,
            )
            for bond in bonds:
                BondImage.objects.create(sub_kyc=subkyc, image=bond)

            messages.success(request, "Sub-KYC submitted successfully.")
        else:
            mykyc = MyKYC.objects.create(
                created_by=user,
                name=name,
                age=age,
                fathername=fathername,
                mobile_number=mobile,
                aadhar_number=aadhar,
                aadhar_front_image=aadhar_front_image,
                aadhar_back_image=aadhar_back_image,
                address=address,
                profession=profession,
                contactSH=contactSH,
                nameSH=nameSH,
                investmentamt=investmentamt,
                passportphoto=passportphoto,
            )
            for bond in bonds:
                BondImage.objects.create(my_kyc=mykyc, image=bond)

            messages.success(request, "My KYC submitted successfully.")

        return redirect("formpage")

    return render(request, "formpage.html", {
        "users": users,
        "my_kyc_list": my_kyc_list,
        "sub_kyc_list": sub_kyc_list,
        "is_sub_mainuser": user.is_sub_mainuser,
        "is_main_user": user.is_main_user,
        "selected_user": selected_user,
    })



from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import MyKYC, SubKYC, BondImage

@login_required
def edit_kyc(request, kyc_id, kyc_type):
    if kyc_type == "my":
        kyc = get_object_or_404(MyKYC, id=kyc_id)
    else:
        kyc = get_object_or_404(SubKYC, id=kyc_id)

    if not (request.user.is_main_user or request.user == kyc.created_by):
        messages.error(request, "You are not authorized to edit this entry.")
        return redirect("formpage")

    if request.method == "POST":
        kyc.name = request.POST.get("name")
        kyc.age = request.POST.get("age")
        kyc.fathername = request.POST.get("fathername")
        kyc.mobile_number = request.POST.get("mobile_number")
        kyc.aadhar_number = request.POST.get("aadhar_number")
        kyc.address = request.POST.get("address")
        kyc.profession = request.POST.get("profession")
        kyc.contactSH = request.POST.get("contactSH")
        kyc.nameSH = request.POST.get("nameSH")
        kyc.investmentamt = request.POST.get("investmentamt")

        if request.FILES.get("aadhar_front_image"):
            kyc.aadhar_front_image = request.FILES["aadhar_front_image"]
        if request.FILES.get("aadhar_back_image"):
            kyc.aadhar_back_image = request.FILES["aadhar_back_image"]
        if request.FILES.get("passportphoto"):
            kyc.passportphoto = request.FILES["passportphoto"]

        # Update existing bond
        bond_ids = request.POST.getlist("bond_id")
        for bond_id in bond_ids:
            try:
                bond = BondImage.objects.get(id=bond_id)
            except BondImage.DoesNotExist:
                continue

            # Update image if new one is uploaded
            new_image = request.FILES.get(f"bond_image_{bond_id}")
            if new_image:
                bond.image = new_image

            bond.companyname = request.POST.get(f"companyname_{bond_id}", "")
            bond.projectname = request.POST.get(f"projectname_{bond_id}", "")
            bond.amount = request.POST.get(f"amount_{bond_id}") or 0
            bond.investment_date = request.POST.get(f"investment_date_{bond_id}") or None
            bond.customer_id = request.POST.get(f"customer_id_{bond_id}", "")
            bond.save()

            # Delete selected bond images
            delete_ids = request.POST.getlist("delete_bonds")
            for bond_id in delete_ids:
                bond = BondImage.objects.filter(id=bond_id).first()
                if bond:
                    bond.image.delete(save=False)
                    bond.delete()

            # Add new bond images
        new_bond_files = request.FILES.getlist("bonds")
        new_company_names = request.POST.getlist("new_companyname")
        new_project_names = request.POST.getlist("new_projectname")
        new_amounts = request.POST.getlist("new_amount")
        new_dates = request.POST.getlist("new_investment_date")
        new_customer_ids = request.POST.getlist("new_customer_id")

        for i, bond_img in enumerate(new_bond_files):
            bond_data = {
                "image": bond_img,
                "companyname": new_company_names[i] if i < len(new_company_names) else "",
                "projectname": new_project_names[i] if i < len(new_project_names) else "",
                "amount": new_amounts[i] if i < len(new_amounts) else 0,
                "investment_date": new_dates[i] if i < len(new_dates) else None,
                "customer_id": new_customer_ids[i] if i < len(new_customer_ids) else "",
            }

            if kyc_type == "my":
                BondImage.objects.create(my_kyc=kyc, **bond_data)
            else:
                BondImage.objects.create(sub_kyc=kyc, **bond_data)

        kyc.save()
        messages.success(request, "KYC updated successfully.")
        return redirect("formpage")

    return render(request, "edit_kyc.html", {"kyc": kyc, "kyc_type": kyc_type})


# views.py
from django.shortcuts import get_object_or_404, redirect
from django.contrib import messages
from .models import MyKYC, SubKYC
from django.contrib.auth.decorators import login_required

@login_required
def delete_kyc(request, kyc_id, kyc_type):
    if kyc_type == "my":
        kyc = get_object_or_404(MyKYC, id=kyc_id)
        key = "hidden_my_kyc"
    elif kyc_type == "sub":
        kyc = get_object_or_404(SubKYC, id=kyc_id)
        key = "hidden_sub_kyc"
    else:
        messages.error(request, "Invalid KYC type.")
        return redirect("formpage")

    if not (request.user.is_main_user or request.user == kyc.created_by):
        messages.error(request, "You are not authorized to delete this entry.")
        return redirect("formpage")

    if request.user.is_main_user:
        kyc.delete()
        messages.success(request, "KYC permanently deleted.")
    else:
        # Track deleted IDs in session
        hidden_ids = request.session.get(key, [])
        if kyc_id not in hidden_ids:
            hidden_ids.append(kyc_id)
        request.session[key] = hidden_ids
        messages.success(request, "KYC removed from your view.")

    return redirect("formpage")


import openpyxl
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from .models import MyKYC, SubKYC

@login_required
def download_kyc_excel(request, kyc_type):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "KYC Details"

    # Headers
    headers = [
        'S.No', 'Name', 'Age', 'Father Name', 'Mobile Number', 'Aadhar Number',
        'Address', 'Profession', 'Contact SH/NH', 'Name SH/NH', 'Investment Amount'
    ]
    sheet.append(headers)

    user = request.user
    selected_user = user

    # Get user_id from query param
    user_id = request.GET.get('user_id')
    if user.is_main_user and user_id:
        from django.contrib.auth import get_user_model
        User = get_user_model()
        try:
            selected_user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            selected_user = user  # fallback

    hidden_ids = request.session.get('hidden_my_kyc' if kyc_type == 'my' else 'hidden_sub_kyc', [])

    if kyc_type == 'my':
        kyc_list = MyKYC.objects.filter(created_by=selected_user)
    else:
        kyc_list = SubKYC.objects.filter(
            user=selected_user
        ) | SubKYC.objects.filter(
            created_by=selected_user
        )
        kyc_list = kyc_list.distinct()

    kyc_list = kyc_list.exclude(id__in=hidden_ids)

    for idx, kyc in enumerate(kyc_list, start=1):
        sheet.append([
            idx,
            kyc.name,
            kyc.age,
            kyc.fathername,
            kyc.mobile_number,
            kyc.aadhar_number,
            kyc.address,
            kyc.profession,
            kyc.contactSH,
            kyc.nameSH,
            kyc.investmentamt
        ])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = 'my_kyc_details.xlsx' if kyc_type == 'my' else 'sub_kyc_details.xlsx'
    response['Content-Disposition'] = f'attachment; filename={filename}'
    workbook.save(response)
    return response


# pdf download

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.units import mm
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from .models import MyKYC, SubKYC

@login_required
def download_kyc_pdf(request, kyc_type):
    response = HttpResponse(content_type='application/pdf')
    filename = "my_kyc_report.pdf" if kyc_type == 'my' else "sub_kyc_report.pdf"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    c = canvas.Canvas(response, pagesize=A4)
    W, H = A4
    margin = 20 * mm
    card_width = W - 2 * margin
    card_height = 60 * mm
    x0 = margin
    y = H - margin

    c.setFont("Helvetica-Bold", 18)
    c.drawCentredString(W / 2, y, "KYC Report")
    y -= 15 * mm

    user = request.user
    selected_user = user

    user_id = request.GET.get('user_id')
    if user.is_main_user and user_id:
        from django.contrib.auth import get_user_model
        User = get_user_model()
        try:
            selected_user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            selected_user = user

    hidden_ids = request.session.get('hidden_my_kyc' if kyc_type == 'my' else 'hidden_sub_kyc', [])

    if kyc_type == 'my':
        kyc_list = MyKYC.objects.filter(created_by=selected_user)
    else:
        kyc_list = SubKYC.objects.filter(
            user=selected_user
        ) | SubKYC.objects.filter(
            created_by=selected_user
        )
        kyc_list = kyc_list.distinct()

    kyc_list = kyc_list.exclude(id__in=hidden_ids)

    for idx, kyc in enumerate(kyc_list, 1):
        if y - card_height < margin:
            c.showPage()
            y = H - margin
            c.setFont("Helvetica-Bold", 18)
            c.drawCentredString(W / 2, y, "KYC Report")
            y -= 15 * mm

        c.setLineWidth(1)
        c.roundRect(x0, y - card_height, card_width, card_height, 5 * mm, stroke=1, fill=0)

        header_h = 10 * mm
        c.setFillColor(colors.lightgrey)
        c.roundRect(x0, y - header_h, card_width, header_h, 5 * mm, stroke=0, fill=1)
        c.setFillColor(colors.black)
        c.setFont("Helvetica-Bold", 12)
        c.drawString(x0 + 5 * mm, y - header_h + 2 * mm, f"KYC #{idx}")

        labels = [
            ("Name", kyc.name),
            ("Father's Name", kyc.fathername or "—"),
            ("Mobile", kyc.mobile_number),
            ("Aadhar", kyc.aadhar_number),
            ("Address", kyc.address),
            ("Profession", kyc.profession or "—"),
            ("Contact SH", kyc.contactSH or "—"),
            ("Name SH", kyc.nameSH or "—"),
            ("Investment", str(kyc.investmentamt) if kyc.investmentamt else "—"),
        ]

        col_x = [x0 + 5 * mm, x0 + card_width / 2 + 5 * mm]
        c.setFont("Helvetica", 10)
        line_h = 6 * mm
        start_y = y - header_h - 5 * mm

        for i, (label, val) in enumerate(labels):
            col = i % 2
            row = i // 2
            text_y = start_y - row * line_h
            c.drawString(col_x[col], text_y, f"{label}: {val}")

        y -= card_height + 5 * mm

    c.save()
    return response




# from reportlab.lib.pagesizes import A4
# from reportlab.pdfgen import canvas
# from reportlab.lib import colors
# from reportlab.lib.units import mm
# from django.http import HttpResponse
# from django.contrib.auth.decorators import login_required
# from .models import MyKYC, SubKYC

# @login_required
# def download_kyc_pdf(request, kyc_type):
#     response = HttpResponse(content_type='application/pdf')
#     filename = "my_kyc_report.pdf" if kyc_type == 'my' else "sub_kyc_report.pdf"
#     response['Content-Disposition'] = f'attachment; filename="{filename}"'

#     c = canvas.Canvas(response, pagesize=A4)
#     W, H = A4
#     margin = 20 * mm
#     card_width = W - 2 * margin
#     card_height = 60 * mm
#     x0 = margin
#     y = H - margin

#     c.setFont("Helvetica-Bold", 18)
#     c.drawCentredString(W / 2, y, "KYC Report")
#     y -= 15 * mm

#     # Select model based on kyc_type
#     user = request.user
#     if kyc_type == 'my':
#         if user.is_superuser or getattr(user, "is_main_user", False):
#             kyc_list = MyKYC.objects.all()
#         elif getattr(user, "is_sub_mainuser", False):
#             kyc_list = MyKYC.objects.filter(user__parent=user)
#         else:
#             kyc_list = MyKYC.objects.filter(user=user)
#     else:
#         if user.is_superuser or getattr(user, "is_main_user", False):
#             kyc_list = SubKYC.objects.all()
#         elif getattr(user, "is_sub_mainuser", False):
#             kyc_list = SubKYC.objects.filter(user__parent=user)
#         else:
#             kyc_list = SubKYC.objects.filter(created_by=user).exclude(user=user)

#     for idx, kyc in enumerate(kyc_list, 1):
#         if y - card_height < margin:
#             c.showPage()
#             y = H - margin
#             c.setFont("Helvetica-Bold", 18)
#             c.drawCentredString(W / 2, y, "KYC Report")
#             y -= 15 * mm

#         c.setLineWidth(1)
#         c.roundRect(x0, y - card_height, card_width, card_height, 5 * mm, stroke=1, fill=0)

#         header_h = 10 * mm
#         c.setFillColor(colors.lightgrey)
#         c.roundRect(x0, y - header_h, card_width, header_h, 5 * mm, stroke=0, fill=1)
#         c.setFillColor(colors.black)
#         c.setFont("Helvetica-Bold", 12)
#         c.drawString(x0 + 5 * mm, y - header_h + 2 * mm, f"KYC #{idx}")

#         labels = [
#             ("Name", kyc.name),
#             ("Father's Name", kyc.fathername or "—"),
#             ("Mobile", kyc.mobile_number),
#             ("Aadhar", kyc.aadhar_number),
#             ("Address", kyc.address),
#             ("Profession", kyc.profession or "—"),
#             ("Contact SH", kyc.contactSH or "—"),
#             ("Name SH", kyc.nameSH or "—"),
#             ("Investment", str(kyc.investmentamt) if kyc.investmentamt else "—"),
#         ]

#         col_x = [x0 + 5 * mm, x0 + card_width / 2 + 5 * mm]
#         c.setFont("Helvetica", 10)
#         line_h = 6 * mm
#         start_y = y - header_h - 5 * mm

#         for i, (label, val) in enumerate(labels):
#             col = i % 2
#             row = i // 2
#             text_y = start_y - row * line_h
#             c.drawString(col_x[col], text_y, f"{label}: {val}")

#         y -= card_height + 5 * mm

#     c.save()
#     return response



################################################################
# subkyc excel and pdf

from django.shortcuts import render
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import MyKYC, SubKYC, BondImage
from django.contrib.auth import get_user_model

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.units import mm

import openpyxl
from openpyxl.styles import Font, Alignment

User = get_user_model()


@login_required
def download_subkyc_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="sub_kyc_report.pdf"'

    c = canvas.Canvas(response, pagesize=A4)
    W, H = A4
    margin = 20 * mm
    card_width = W - 2 * margin
    card_height = 60 * mm
    x0 = margin
    y = H - margin

    c.setFont("Helvetica-Bold", 18)
    c.drawCentredString(W / 2, y, "Sub-KYC Report")
    y -= 15 * mm

    user = request.user
    hidden_sub_kyc_ids = request.session.get('hidden_sub_kyc', [])

    if user.is_superuser or user.is_main_user:
        kyc_list = SubKYC.objects.filter(created_by=user) | SubKYC.objects.filter(user=user)
    else:
        kyc_list = SubKYC.objects.filter(user=user)

    kyc_list = kyc_list.exclude(id__in=hidden_sub_kyc_ids).distinct()

    for idx, kyc in enumerate(kyc_list, 1):
        if y - card_height < margin:
            c.showPage()
            y = H - margin
            c.setFont("Helvetica-Bold", 18)
            c.drawCentredString(W / 2, y, "Sub-KYC Report")
            y -= 15 * mm

        c.setLineWidth(1)
        c.roundRect(x0, y - card_height, card_width, card_height, 5 * mm, stroke=1, fill=0)

        header_h = 10 * mm
        c.setFillColor(colors.lightgrey)
        c.roundRect(x0, y - header_h, card_width, header_h, 5 * mm, stroke=0, fill=1)
        c.setFillColor(colors.black)
        c.setFont("Helvetica-Bold", 12)
        c.drawString(x0 + 5 * mm, y - header_h + 2 * mm, f"Sub-KYC #{idx}")

        labels = [
            ("Name", kyc.name),
            ("Father's Name", kyc.fathername or "—"),
            ("Mobile", kyc.mobile_number),
            ("Aadhar", kyc.aadhar_number),
            ("Address", kyc.address),
            ("Profession", kyc.profession or "—"),
            ("Contact SH", kyc.contactSH or "—"),
            ("Name SH", kyc.nameSH or "—"),
            ("Investment", str(kyc.investmentamt) if kyc.investmentamt else "—"),
        ]

        col_x = [x0 + 5 * mm, x0 + card_width / 2 + 5 * mm]
        c.setFont("Helvetica", 10)
        line_h = 6 * mm
        start_y = y - header_h - 5 * mm

        for i, (label, val) in enumerate(labels):
            col = i % 2
            row = i // 2
            text_y = start_y - row * line_h
            c.drawString(col_x[col], text_y, f"{label}: {val}")

        y -= card_height + 5 * mm

    c.save()
    return response



@login_required
def download_subkyc_excel(request):
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="sub_kyc_report.xlsx"'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sub-KYC Report"

    headers = [
        "Name", "Father's Name", "Mobile", "Aadhar",
        "Address", "Profession", "Contact SH",
        "Name SH", "Investment"
    ]
    ws.append(headers)

    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    user = request.user
    hidden_sub_kyc_ids = request.session.get('hidden_sub_kyc', [])

    if user.is_superuser or user.is_main_user:
        kyc_list = SubKYC.objects.filter(created_by=user) | SubKYC.objects.filter(user=user)
    else:
        kyc_list = SubKYC.objects.filter(user=user)

    kyc_list = kyc_list.exclude(id__in=hidden_sub_kyc_ids).distinct()

    for kyc in kyc_list:
        ws.append([
            kyc.name,
            kyc.fathername or "",
            kyc.mobile_number,
            kyc.aadhar_number,
            kyc.address,
            kyc.profession or "",
            kyc.contactSH or "",
            kyc.nameSH or "",
            kyc.investmentamt if kyc.investmentamt else 0,
        ])

    wb.save(response)
    return response



@login_required
def profile_view(request):
    user = request.user

    if request.method == 'POST':
        user.username = request.POST['username']
        user.email = request.POST['email']
        user.phone_number = request.POST['phone_number']
        user.save()
        messages.success(request, "Profile updated successfully.")
        return redirect('profile')

    return render(request, 'profile.html', {'user_obj': user})



from django.contrib.auth.decorators import login_required
from django.shortcuts import redirect
from django.contrib import messages
from .models import MyKYC, SubKYC, BondImage, User

@login_required
def add_my_kyc(request):
    if request.method == "POST":
        created_by = request.user

        try:
            data_for_user = User.objects.get(id=request.POST.get('data_for_user'))
        except User.DoesNotExist:
            messages.error(request, "Invalid user selected.")
            return redirect("formpage")    

        try:
            investmentamt = int(request.POST.get('investmentamt')) if request.POST.get('investmentamt') else None
        except ValueError:
            messages.error(request, "Investment amount must be numeric.")
            return redirect('formpage')

        my_kyc = MyKYC.objects.create(
            user=data_for_user,
            created_by=created_by,
            name=request.POST.get('name'),
            age=request.POST.get('age'),
            fathername=request.POST.get('fathername'),
            mobile_number=request.POST.get('mobile_number'),
            aadhar_number=request.POST.get('aadhar_number'),
            aadhar_front_image=request.FILES.get('aadhar_front_image'),
            aadhar_back_image=request.FILES.get('aadhar_back_image'),
            address=request.POST.get('address'),
            profession=request.POST.get('profession'),
            contactSH=request.POST.get('contactSH'),
            nameSH=request.POST.get('nameSH'),
            investmentamt=investmentamt,
            passportphoto=request.FILES.get('passportphoto'),
        )

        # for bond_file in request.FILES.getlist('bonds'):
        #     BondImage.objects.create(my_kyc=my_kyc, image=bond_file)

        bond_files = request.FILES.getlist('bonds')

        for i, bond_file in enumerate(bond_files):
            BondImage.objects.create(
                my_kyc=my_kyc,
                image=bond_file,
                companyname=request.POST.get(f'companyname_{i}'),
                projectname=request.POST.get(f'projectname_{i}'),
                amount=request.POST.get(f'amount_{i}') or 0,
                investment_date=request.POST.get(f'investment_date_{i}') or datetime.date.today(),
                customer_id=request.POST.get(f'customer_id_{i}')
            )

        messages.success(request, "My KYC data saved successfully.")
        return redirect('formpage')



from django.contrib.auth.decorators import login_required
from django.shortcuts import redirect
from django.contrib import messages
from .models import SubKYC, BondImage, User

@login_required
def add_other_kyc(request):
    if request.method == "POST":
        created_by = request.user

        try:
            data_for_user = User.objects.get(id=request.POST.get('data_for_user'))
        except User.DoesNotExist:
            messages.error(request, "Invalid user selected.")
            return redirect("formpage")

        try:
            investmentamt = int(request.POST.get('investmentamt')) if request.POST.get('investmentamt') else None
        except ValueError:
            messages.error(request, "Investment amount must be numeric.")
            return redirect("formpage")

        kyc = SubKYC.objects.create(
            user=data_for_user,
            created_by=created_by,
            name=request.POST.get('name'),
            age=request.POST.get('age'),
            mobile_number=request.POST.get('mobile_number'),
            fathername=request.POST.get('fathername'),
            address=request.POST.get('address'),
            aadhar_number=request.POST.get('aadhar_number'),
            profession=request.POST.get('profession'),
            contactSH=request.POST.get('contactSH'),
            nameSH=request.POST.get('nameSH'),
            investmentamt=investmentamt,
            aadhar_front_image=request.FILES.get('aadhar_front_image'),
            aadhar_back_image=request.FILES.get('aadhar_back_image'),
            passportphoto=request.FILES.get('passportphoto'),
        )

        for bond_file in request.FILES.getlist('bonds'):
            BondImage.objects.create(sub_kyc=kyc, image=bond_file)  # ✅ FIXED HERE

        messages.success(request, "Sub-KYC added successfully.")
        return redirect('formpage')


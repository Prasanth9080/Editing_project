# views.py
import datetime
from django.shortcuts import render, redirect
from django.contrib import messages
from django.contrib.auth import login, logout
from rest_framework_simplejwt.tokens import RefreshToken
from .models import User
import random, string

otp_storage = {}

def generate_otp():
    return ''.join(random.choices(string.digits, k=6))  # 6-digit OTP


# from django.shortcuts import render, redirect
# from django.contrib import messages
# from django.contrib.auth.models import BaseUserManager
# from rest_framework_simplejwt.tokens import RefreshToken

# from .models import User  # your custom user model

# def signup_view(request):
#     if request.method == 'POST':
#         name = request.POST.get('username')
#         phone = request.POST.get('phone_number')
#         email = request.POST.get('email')

#         # Server-side validation
#         if len(phone) != 10 or not phone.isdigit():
#             messages.error(request, "Enter a valid 10-digit phone number")
#             return redirect('signup')

#         if User.objects.filter(phone_number=phone).exists():
#             messages.error(request, "Phone number already registered")
#             return redirect('signup')

#         if User.objects.filter(email=email).exists():
#             messages.error(request, "Email already registered")
#             return redirect('signup')

#         # Generate random password using BaseUserManager
#         # random_password = BaseUserManager().make_random_password()

#         # Create user
#         user = User.objects.create_user(
#             username=name,
#             email=email,
#             phone_number=phone,
#             # password=random_password
#         )

#         # Generate JWT token
#         refresh = RefreshToken.for_user(user)
#         user.jwt_token = str(refresh.access_token)
#         user.save()

#         messages.success(request, "Signup successful! Please login.")
#         return redirect('login')

#     return render(request, 'signup.html')


# Signup

def signup_view(request):
    if request.method == 'POST':
        name = request.POST.get('username')
        phone = request.POST.get('phone_number')
        email = request.POST.get('email')
        password = request.POST.get('password')

        if " " in name:
            messages.error(request, "Username must not contain spaces")
            return redirect('signup')

        if User.objects.filter(username=name).exists():
            messages.error(request, "Username is already exists")
            return redirect('signup')

        if len(phone) != 10 or not phone.isdigit():
            messages.error(request, "Enter a valid 10-digit phone number")
            return redirect('signup')

        if User.objects.filter(phone_number=phone).exists():
            messages.error(request, "Phone number already registered")
            return redirect('signup')

        if User.objects.filter(email=email).exists():
            messages.error(request, "Email already registered")
            return redirect('signup')

        user = User.objects.create_user(
            username=name,
            email=email, 
            phone_number=phone,
            password=password
        )

        refresh = RefreshToken.for_user(user)
        user.jwt_token = str(refresh.access_token)
        user.save()

        messages.success(request, "Signup successful! Please login.")
        return redirect('login')

    return render(request, 'signup.html')

# change password

from django.contrib.auth.decorators import login_required
from django.contrib.auth import update_session_auth_hash
from django.contrib import messages
from django.shortcuts import render, redirect

@login_required
def change_password_view(request):
    if request.method == 'POST':
        old_password = request.POST.get('old_password')
        new_password1 = request.POST.get('new_password1')
        new_password2 = request.POST.get('new_password2')

        user = request.user

        if not user.check_password(old_password):
            messages.error(request, '❌ Old password is incorrect.')
        elif new_password1 != new_password2:
            messages.error(request, '❌ New passwords do not match.')
        else:
            user.set_password(new_password1)
            user.save()
            update_session_auth_hash(request, user)  # prevent logout
            messages.success(request, '✅ Password changed successfully.')
            return redirect('profile')  # change this to your success page

    return render(request, 'change_password.html')

# test-mail

# from django.core.mail import send_mail
# from django.http import HttpResponse

# def test_email(request):
#     send_mail(
#         'Test Subject',
#         'This is a test email from Django.',
#         'prasanthchaandhu02@gmail.com',  # must match EMAIL_HOST_USER in settings
#         ['prasanthchaandhu02@gmail.com'],  # recipient email
#         fail_silently=False,
#     )
#     return HttpResponse("Email sent!")

# password reset through email

# views.py
# views.py
from django.shortcuts import render, redirect
from .forms import CustomPasswordResetForm
from django.contrib import messages
from .forms import CustomPasswordResetForm
from django.contrib import messages

def password_reset_request(request):
    form = CustomPasswordResetForm(request.POST or None)
    form = CustomPasswordResetForm(request.POST or None)
    if request.method == "POST":
        if form.is_valid():
            form.save(
                request=request,
                use_https=False,
                # use_https=False,
                from_email='prasanthchaandhu02@gmail.com',
                email_template_name='password_reset_email.html',
            )
            messages.success(request, "Password reset link sent! Please check your email.")
            messages.success(request, "Password reset link sent! Please check your email.")
            return redirect('password_reset_done')
        else:
            messages.error(request, "Email is not valid. Please enter a valid email address.")
    return render(request, 'password_reset.html', {'form': form})


# def login_view(request):
#     if request.method == 'POST':
#         phone = request.POST.get('phone_number')

#         try:
#             user = User.objects.get(phone_number=phone)

#             if not user.jwt_token:
#                 # Generate new token if missing
#                 refresh = RefreshToken.for_user(user)
#                 user.jwt_token = str(refresh.access_token)
#                 user.save()

#             otp = generate_otp()
#             otp_storage[phone] = otp
#             request.session['phone_number'] = phone

#             print(f"\n🔐 OTP for {phone}: {otp}\n")  # Print clearly to terminal

#             messages.success(request, "OTP sent to your number.")
#             return redirect('verify_otp')

#         except User.DoesNotExist:
#             messages.error(request, "This number is not registered. Please sign up.")
#             return redirect('signup')

#     return render(request, 'login.html')

# Login 

from django.shortcuts import render, redirect
from django.contrib.auth import login
from django.contrib import messages
from .models import User
 
def login_view(request):
    if request.method == 'POST':
        phone = request.POST.get('phone_number')
        password = request.POST.get('password')

        try:
            user = User.objects.get(phone_number=phone)

            if user.check_password(password):
                login(request, user)
                messages.success(request, "Login successful.")
                return redirect('formpage')
            else:
                messages.error(request, "Incorrect password.")
                return redirect('login')

        except User.DoesNotExist:
            messages.error(request, "User with this phone number does not exist.")
            return redirect('login')

    return render(request, 'login.html')


# Phone number

from django.contrib.auth.backends import ModelBackend
from .models import User

class PhoneNumberBackend(ModelBackend):
    def authenticate(self, request, phone_number=None, password=None, **kwargs):
        try:
            user = User.objects.get(phone_number=phone_number)
            if user.check_password(password):
                return user
        except User.DoesNotExist:
            return None


def verify_otp_view(request):
    phone = request.session.get('phone_number')
    if not phone:
        return redirect('login')

    if request.method == 'POST':
        input_otp = request.POST.get('otp')
        expected_otp = otp_storage.get(phone)

        if input_otp == expected_otp:
            try:
                user = User.objects.get(phone_number=phone)
            except User.DoesNotExist:
                messages.error(request, "User not found.")
                return redirect('login')

            login(request, user)

            # Generate access token
            refresh = RefreshToken.for_user(user)
            access_token = str(refresh.access_token)
            request.session['access_token'] = access_token

            response = redirect('formpage')
            response.set_cookie('auth_token', access_token)
            return response
        else:
            messages.error(request, "Invalid OTP.")

    return render(request, 'verify_otp.html')


def logout_view(request):
    logout(request)
    messages.success(request, "Logged out successfully.")
    response = redirect('login')
    response.delete_cookie('jwt_token')
    return response



################ delete option only acccess in main user
################ otherwise normal user delete the record only delete(hide) the paricular role
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import MyKYC, SubKYC, BondImage
from django.contrib.auth import get_user_model

User = get_user_model()

@login_required
def form_page(request):
    user = request.user
    user_id = request.GET.get('user_id')
    selected_user = None

    hidden_my_kyc_ids = request.session.get('hidden_my_kyc', [])
    hidden_sub_kyc_ids = request.session.get('hidden_sub_kyc', [])

    # Determine selected user
    if user.is_main_user or user.is_sub_mainuser:
        if user_id:
            try:
                selected_user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                messages.error(request, "Selected user does not exist.")
                return redirect("formpage")
        else:
            selected_user = user
    else:
        selected_user = user

    # Query MyKYC & SubKYC
    my_kyc_list = MyKYC.objects.filter(created_by=selected_user).exclude(id__in=hidden_my_kyc_ids)
    if user.is_main_user or user.is_sub_mainuser:
        sub_kyc_list = (
            SubKYC.objects.filter(user=selected_user) |
            SubKYC.objects.filter(created_by=selected_user)
        ).exclude(id__in=hidden_sub_kyc_ids).distinct()
    else:
        sub_kyc_list = SubKYC.objects.filter(user=user).exclude(id__in=hidden_sub_kyc_ids)

    users = User.objects.filter(is_superuser=False, is_main_user=False, is_sub_mainuser=False).exclude(id=user.id)

    # Handle form submission
    if request.method == "POST":
        # Basic KYC fields
        kyc_data = {
            "membershipno": request.POST.get("membershipno"),
            "membershiptype": request.POST.get("membershiptype"),
            "depositorsname": request.POST.get("depositorsname"),
            "depositorsaddress": request.POST.get("depositorsaddress"),
            "bondholdername": request.POST.get("bondholdername"),
            # "projectname": request.POST.get("projectname"),
            "depositormobile_number": request.POST.get("depositormobile_number"),
            "agentname": request.POST.get("agentname"),
            "agentaddress": request.POST.get("agentaddress"),
            "nameofdirector": request.POST.get("nameofdirector"),
            "aadhar_number": request.POST.get("aadhar_number"),
            "pan_number": request.POST.get("pan_number"),
            "ration_number": request.POST.get("ration_number"),
            "bankname": request.POST.get("bankname"),
            "bankaccno": request.POST.get("bankaccno"),
            "ifscno": request.POST.get("ifscno"),
        }

        # File uploads
        kyc_data["aadhar_front_image"] = request.FILES.get("aadhar_front_image")
        kyc_data["aadhar_back_image"] = request.FILES.get("aadhar_back_image")
        kyc_data["passportphoto"] = request.FILES.get("passportphoto")

        # Bond details (dynamic fields)
        bond_entries = []
        index = 0
        while True:
            company_key = f"companyname_{index}"
            if company_key not in request.POST:
                break
            bond_entries.append({
                ""
                "bondholdername": request.POST.get(f"bondholdername_{index}"),
                "bond_image": request.FILES.getlist("bonds")[index] if len(request.FILES.getlist("bonds")) > index else None,
                "company_name": request.POST.get(company_key),
                "project_name": request.POST.get(f"projectname_{index}"),
                "refundamount": request.POST.get(f"refundamount_{index}"),
                "balanceamount": request.POST.get(f"balanceamount_{index}"),
                "amount": request.POST.get(f"amount_{index}"),
                "investment_date": request.POST.get(f"investment_date_{index}"),
                "customer_id": request.POST.get(f"customer_id_{index}"),
            })
            index += 1

        # Determine KYC type
        is_sub_kyc = 'is_sub_kyc' in request.POST
        data_for_user = user
        if is_sub_kyc and (user.is_main_user or user.is_sub_mainuser):
            try:
                data_for_user = User.objects.get(id=request.POST.get("data_for_user"))
            except User.DoesNotExist:
                messages.error(request, "Invalid user selected.")
                return redirect("formpage")

        # Validation (basic example)
        if not kyc_data["depositorsname"] or not kyc_data["aadhar_number"]:
            messages.error(request, "Depositor's name and Aadhar number are required.")
            return redirect("formpage")

        # Save KYC
        if is_sub_kyc:
            kyc_obj = SubKYC.objects.create(created_by=user, user=data_for_user, **kyc_data)
        else:
            kyc_obj = MyKYC.objects.create(created_by=user, **kyc_data)

        # Save bonds
        for bond in bond_entries:
            BondImage.objects.create(
                sub_kyc=kyc_obj if is_sub_kyc else None,
                my_kyc=kyc_obj if not is_sub_kyc else None,
                bondholdername=bond["bondholdername"],
                image=bond["bond_image"],
                company_name=bond["company_name"],
                project_name=bond["project_name"],
                refundamount=bond["refundamount"],
                balanceamount=bond["balanceamount"],
                amount=bond["amount"],
                investment_date=bond["investment_date"],
                customer_id=bond["customer_id"],
            )

        messages.success(request, "KYC submitted successfully.")
        return redirect("formpage")

    return render(request, "formpage.html", {
        "users": users,
        "my_kyc_list": my_kyc_list,
        "sub_kyc_list": sub_kyc_list,
        "is_sub_mainuser": user.is_sub_mainuser,
        "is_main_user": user.is_main_user,
        "selected_user": selected_user,
    })




from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import MyKYC, SubKYC, BondImage

@login_required
def edit_kyc(request, kyc_id, kyc_type):
    if kyc_type == "my":
        kyc = get_object_or_404(MyKYC, id=kyc_id)
    else:
        kyc = get_object_or_404(SubKYC, id=kyc_id)

    if not (request.user.is_main_user or request.user == kyc.created_by):
        messages.error(request, "You are not authorized to edit this entry.")
        return redirect("formpage")

    if request.method == "POST":
        # Text fields
        kyc.membershipno = request.POST.get("membershipno")
        # kyc.membershiptype = request.POST.get("membershiptype")
        kyc.depositorsname = request.POST.get("depositorsname")
        kyc.depositorsaddress = request.POST.get("depositorsaddress")
        # kyc.nameofthecompany = request.POST.get("nameofthecompany")
        # kyc.customeridno = request.POST.get("customeridno")
        # kyc.receiptno = request.POST.get("receiptno")
        # kyc.modno = request.POST.get("modno")
        # kyc.depositamount = request.POST.get("depositamount")
        # kyc.intrefundamount = request.POST.get("intrefundamount")
        # kyc.defaultamount = request.POST.get("defaultamount")
        # kyc.investmentdate = request.POST.get("investmentdate")
        kyc.bondholdername = request.POST.get("bondholdername")
        # kyc.projectname = request.POST.get("projectname")
        kyc.depositormobile_number = request.POST.get("depositormobile_number")
        kyc.aadhar_number = request.POST.get("aadhar_number")
        kyc.pan_number = request.POST.get("pan_number")
        kyc.ration_number = request.POST.get("ration_number")
        kyc.bankname = request.POST.get("bankname")
        kyc.bankaccno = request.POST.get("bankaccno")
        kyc.ifscno = request.POST.get("ifscno")
        kyc.agentname = request.POST.get("agentname")
        kyc.agentaddress = request.POST.get("agentaddress")
        kyc.nameofdirector = request.POST.get("nameofdirector")

        # File fields — update only if a new file is uploaded
        if request.FILES.get("aadhar_front_image"):
            kyc.aadhar_front_image = request.FILES["aadhar_front_image"]
        if request.FILES.get("aadhar_back_image"):
            kyc.aadhar_back_image = request.FILES["aadhar_back_image"]
        if request.FILES.get("passportphoto"):
            kyc.passportphoto = request.FILES["passportphoto"]

        # Update existing bonds
        bond_ids = request.POST.getlist("bond_id")
        for bond_id in bond_ids:
            try:
                bond = BondImage.objects.get(id=bond_id)
            except BondImage.DoesNotExist:
                continue

            # Update bond image if new one is uploaded
            bond.bondholdername = request.POST.get(f"bondholdername_{bond_id}", "")
            new_image = request.FILES.get(f"bond_image_{bond_id}")
            if new_image:
                bond.image = new_image

            bond.companyname = request.POST.get(f"companyname_{bond_id}", "")
            bond.projectname = request.POST.get(f"projectname_{bond_id}", "")
            bond.amount = request.POST.get(f"amount_{bond_id}") or 0
            bond.refundamount = request.POST.get(f"refundamount_{bond_id}", "")
            bond.balanceamount = request.POST.get(f"balanceamount_{bond_id}", "")
            bond.investment_date = request.POST.get(f"investment_date_{bond_id}") or None
            bond.customer_id = request.POST.get(f"customer_id_{bond_id}", "")
            bond.save()

        # Delete selected bond images
        delete_ids = request.POST.getlist("delete_bonds")
        for bond_id in delete_ids:
            bond = BondImage.objects.filter(id=bond_id).first()
            if bond:
                if bond.image:
                    bond.image.delete(save=False)
                bond.delete()

        # Add new bond images
        new_bondholder_names = request.POST.getlist("new_bondholdername")
        new_bond_files = request.FILES.getlist("new_bonds")
        new_company_names = request.POST.getlist("new_companyname")
        new_project_names = request.POST.getlist("new_projectname")
        new_amounts = request.POST.getlist("new_amount")
        new_refundamounts = request.POST.getlist("new_refundamount")
        new_balanceamounts = request.POST.getlist("new_balanceamount")
        new_dates = request.POST.getlist("new_investment_date")
        new_customer_ids = request.POST.getlist("new_customer_id")

        for i, bond_img in enumerate(new_bond_files):
            bond_data = {
                "bondholdername": new_bondholder_names[i] if i < len(new_bondholder_names) else "",
                "image": bond_img,
                "companyname": new_company_names[i] if i < len(new_company_names) else "",
                "projectname": new_project_names[i] if i < len(new_project_names) else "",
                "amount": new_amounts[i] if i < len(new_amounts) else 0,
                "refundamount": new_refundamounts[i] if i < len(new_refundamounts) else "",
                "balanceamount": new_balanceamounts[i] if i < len(new_balanceamounts) else "",
                "investment_date": new_dates[i] if i < len(new_dates) else None,
                "customer_id": new_customer_ids[i] if i < len(new_customer_ids) else "",
            }

            if kyc_type == "my":
                BondImage.objects.create(my_kyc=kyc, **bond_data)
            else:
                BondImage.objects.create(sub_kyc=kyc, **bond_data)

        kyc.save()
        messages.success(request, "KYC updated successfully.")
        return redirect("formpage")

    return render(request, "edit_kyc.html", {"kyc": kyc, "kyc_type": kyc_type})



# views.py
from django.shortcuts import get_object_or_404, redirect
from django.contrib import messages
from .models import MyKYC, SubKYC
from django.contrib.auth.decorators import login_required

@login_required
def delete_kyc(request, kyc_id, kyc_type):
    if kyc_type == "my":
        kyc = get_object_or_404(MyKYC, id=kyc_id)
        key = "hidden_my_kyc"
    elif kyc_type == "sub":
        kyc = get_object_or_404(SubKYC, id=kyc_id)
        key = "hidden_sub_kyc"
    else:
        messages.error(request, "Invalid KYC type.")
        return redirect("formpage")

    if not (request.user.is_main_user or request.user == kyc.created_by):
        messages.error(request, "You are not authorized to delete this entry.")
        return redirect("formpage")

    if request.user.is_main_user:
        kyc.delete()
        messages.success(request, "KYC permanently deleted.")
    else:
        # Track deleted IDs in session
        hidden_ids = request.session.get(key, [])
        if kyc_id not in hidden_ids:
            hidden_ids.append(kyc_id)
        request.session[key] = hidden_ids
        messages.success(request, "KYC removed from your view.")

    return redirect("formpage")


import openpyxl
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.utils.dateformat import format as date_format
from django.contrib.auth import get_user_model
from .models import MyKYC, SubKYC


@login_required
def download_kyc_excel(request, kyc_type):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "KYC Details"

    headers = [
        'S.No', 'Membership No', 'Depositor Name', 'Depositor Address',
        'Mobile', 'Aadhar Number',
        'PAN Number', 'Ration Number',
        'Bank Name', 'Bank A/C No', 'IFSC No', 'Agent Name', 'Agent Address', 'Name of Director', 'Passport Photo URL',
        'Aadhar Front Image URL', 'Aadhar Back Image URL',
        'Bond S.No','Bond Holder Name', 'Image URL', 'Company Name', 'Project Name', 'Deposit Amount',
        'Interest Refund Amount', 'Default Amount' , 'Investment Date', 'Customer ID'
    ]
    sheet.append(headers)

    user = request.user
    selected_user = user

    # If main user and viewing a specific sub-user
    user_id = request.GET.get('user_id')
    if getattr(user, "is_main_user", False) and user_id:
        try:
            selected_user = get_user_model().objects.get(id=user_id)
        except get_user_model().DoesNotExist:
            pass

    # Determine hidden IDs
    hidden_ids_key = 'hidden_my_kyc' if kyc_type == 'my' else 'hidden_sub_kyc'
    hidden_ids = request.session.get(hidden_ids_key, [])

    # Get KYC queryset
    if kyc_type == 'my':
        kyc_list = MyKYC.objects.filter(created_by=selected_user)
    else:
        kyc_list = (
            SubKYC.objects.filter(user=selected_user) |
            SubKYC.objects.filter(created_by=selected_user)
        ).distinct()

    kyc_list = kyc_list.exclude(id__in=hidden_ids)

    for idx, kyc in enumerate(kyc_list, start=1):
        # Safe URL fetch
        aadhar_front_url = kyc.aadhar_front_image.url if kyc.aadhar_front_image and kyc.aadhar_front_image.name else ''
        aadhar_back_url = kyc.aadhar_back_image.url if kyc.aadhar_back_image and kyc.aadhar_back_image.name else ''
        passport_url = kyc.passportphoto.url if kyc.passportphoto and kyc.passportphoto.name else ''

        # Base KYC row (no bond details here)
        base_row = [
            idx,
            kyc.membershipno,
            # kyc.membershiptype,
            kyc.depositorsname,
            kyc.depositorsaddress,
            # kyc.projectname,
            kyc.depositormobile_number,
            kyc.aadhar_number,
            kyc.pan_number,
            kyc.ration_number,
            kyc.bankname,
            kyc.bankaccno,
            kyc.ifscno,
            kyc.agentname,
            kyc.agentaddress,
            kyc.nameofdirector,
            passport_url,
            aadhar_front_url,
            aadhar_back_url,
        ]
        sheet.append(base_row + [''] * 7)  # Empty bond columns for KYC row

        # Bond details separately
        bonds = list(kyc.bonds.all())
        for b_idx, bond in enumerate(bonds, start=1):
            bond_image_url = bond.image.url if bond.image and bond.image.name else ''
            bond_row = [''] * 17  # empty KYC columns
            bond_row += [
                b_idx,
                bond.bondholdername,
                bond_image_url,
                bond.companyname,
                bond.projectname,
                bond.amount,
                bond.refundamount,
                bond.balanceamount,
                date_format(bond.investment_date, 'd-m-Y') if bond.investment_date else '',
                bond.customer_id
            ]
            sheet.append(bond_row)

    # Prepare HTTP response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = 'my_kyc_details.xlsx' if kyc_type == 'my' else 'sub_kyc_details.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response



# pdf download

from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
import os
from .models import MyKYC, SubKYC


@login_required
def download_kyc_pdf(request, kyc_type):
    # Prepare HTTP response
    filename = "my_kyc_report.pdf" if kyc_type == 'my' else "sub_kyc_report.pdf"
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    # Create PDF document
    doc = SimpleDocTemplate(response, pagesize=A4,
                            rightMargin=15*mm, leftMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm)

    styles = getSampleStyleSheet()
    elements = []

    # Select user
    user = request.user
    selected_user = user

    user_id = request.GET.get('user_id')
    if getattr(user, "is_main_user", False) and user_id:
        from django.contrib.auth import get_user_model
        User = get_user_model()
        try:
            selected_user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            selected_user = user

    # Hidden IDs from session
    hidden_ids = request.session.get('hidden_my_kyc' if kyc_type == 'my' else 'hidden_sub_kyc', [])

    if kyc_type == 'my':
        kyc_list = MyKYC.objects.filter(created_by=selected_user)
    else:
        kyc_list = SubKYC.objects.filter(user=selected_user) | SubKYC.objects.filter(created_by=selected_user)
        kyc_list = kyc_list.distinct()

    kyc_list = kyc_list.exclude(id__in=hidden_ids)

    # Loop through KYC records
    for idx, kyc in enumerate(kyc_list, start=1):
        # KYC Header
        elements.append(Paragraph(f"<b>Details Row{idx}</b>", styles['Heading2']))
        elements.append(Spacer(1, 4*mm))

        # Main KYC details
        main_table_data = [
            ["Membership No", kyc.membershipno or "—"],
            # ["Membership Type", kyc.membershiptype or "—"],
            ["Depositor Name", kyc.depositorsname or "—"],
            ["Depositor Address", kyc.depositorsaddress or "—"],
            # ["Project Name", kyc.projectname or "—"],
            ["Depositor Mobile", kyc.depositormobile_number or "—"],
            ["Aadhar No", kyc.aadhar_number or "—"],
            ["PAN No", kyc.pan_number or "—"],
            ["Ration Card No", kyc.ration_number or "—"],
            ["Bank Name", kyc.bankname or "—"],
            ["Account No", kyc.bankaccno or "—"],
            ["IFSC Code", kyc.ifscno or "—"],
            ["Agent Name", kyc.agentname or "—"],
            ["Agent Address", kyc.agentaddress or "—"],
            ["Director Name", kyc.nameofdirector or "—"],
            ["Passport Photo", Image(kyc.passportphoto.path, width=20*mm, height=25*mm) if kyc.passportphoto and os.path.exists(kyc.passportphoto.path) else "—"],
            ["Aadhar Front", Image(kyc.aadhar_front_image.path, width=20*mm, height=15*mm) if kyc.aadhar_front_image and os.path.exists(kyc.aadhar_front_image.path) else "—"],
            ["Aadhar Back", Image(kyc.aadhar_back_image.path, width=20*mm, height=15*mm) if kyc.aadhar_back_image and os.path.exists(kyc.aadhar_back_image.path) else "—"],

        ]

        main_table = Table(main_table_data, colWidths=[50*mm, 110*mm])
        main_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTNAME', (1, 0), (1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('ALIGN', (0, 0), (0, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0, 0), (-1, -1), [colors.whitesmoke, colors.lightyellow]),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.black),
            ('INNERGRID', (0, 0), (-1, -1), 0.25, colors.black),
        ]))
        elements.append(main_table)
        elements.append(Spacer(1, 5*mm))

        # Bond Details (Multiple Bonds Support)
        bonds = kyc.bonds.all()
        for b_idx, bond in enumerate(bonds, start=1):
            elements.append(Paragraph(f"<b>Bond Details Row{b_idx}</b>", styles['Heading3']))
            elements.append(Spacer(1, 2*mm))

            bond_table_data = [
                ["Bond Holder Name", bond.bondholdername or "—"],
                ["Investment Date", bond.investment_date.strftime('%d-%m-%Y') if bond.investment_date else "—"],
                ["Bond Image", Image(bond.image.path, width=20*mm, height=15*mm) if bond.image and os.path.exists(bond.image.path) else "—"],
                ["Project Name", bond.projectname or "—"],
                ["Amount", getattr(bond, "amount", "—") or "—"],
                ["refundamount", getattr(bond, "refundamount", "—") or "—"],
                ["balanceamount", getattr(bond, "balanceamount", "—") or "—"],
                ["Company Name", getattr(bond, "companyname", "—") or "—"],
                ["Customer ID", getattr(bond, "customer_id", "—") or "—"],
            ]

            bond_table = Table(bond_table_data, colWidths=[50*mm, 110*mm])
            bond_table.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTNAME', (1, 0), (1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('ALIGN', (0, 0), (0, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ROWBACKGROUNDS', (0, 0), (-1, -1), [colors.whitesmoke, colors.lightyellow]),
                ('BOX', (0, 0), (-1, -1), 0.5, colors.black),
                ('INNERGRID', (0, 0), (-1, -1), 0.25, colors.black),
            ]))
            elements.append(bond_table)
            elements.append(Spacer(1, 5*mm))

        elements.append(Spacer(1, 10*mm))

    # Build PDF
    doc.build(elements)
    return response




# from reportlab.lib.pagesizes import A4
# from reportlab.pdfgen import canvas
# from reportlab.lib import colors
# from reportlab.lib.units import mm
# from django.http import HttpResponse
# from django.contrib.auth.decorators import login_required
# from .models import MyKYC, SubKYC

# @login_required
# def download_kyc_pdf(request, kyc_type):
#     response = HttpResponse(content_type='application/pdf')
#     filename = "my_kyc_report.pdf" if kyc_type == 'my' else "sub_kyc_report.pdf"
#     response['Content-Disposition'] = f'attachment; filename="{filename}"'

#     c = canvas.Canvas(response, pagesize=A4)
#     W, H = A4
#     margin = 20 * mm
#     card_width = W - 2 * margin
#     card_height = 60 * mm
#     x0 = margin
#     y = H - margin

#     c.setFont("Helvetica-Bold", 18)
#     c.drawCentredString(W / 2, y, "KYC Report")
#     y -= 15 * mm

#     user = request.user
#     selected_user = user
 
#     user_id = request.GET.get('user_id')
#     if user.is_main_user and user_id:
#         from django.contrib.auth import get_user_model
#         User = get_user_model()
#         try:
#             selected_user = User.objects.get(id=user_id)
#         except User.DoesNotExist:
#             selected_user = user

#     hidden_ids = request.session.get('hidden_my_kyc' if kyc_type == 'my' else 'hidden_sub_kyc', [])

#     if kyc_type == 'my':
#         kyc_list = MyKYC.objects.filter(created_by=selected_user)
#     else:
#         kyc_list = SubKYC.objects.filter(
#             user=selected_user
#         ) | SubKYC.objects.filter(
#             created_by=selected_user
#         )
#         kyc_list = kyc_list.distinct()

#     kyc_list = kyc_list.exclude(id__in=hidden_ids)

#     for idx, kyc in enumerate(kyc_list, 1):
#         if y - card_height < margin:
#             c.showPage()
#             y = H - margin
#             c.setFont("Helvetica-Bold", 18)
#             c.drawCentredString(W / 2, y, "KYC Report")
#             y -= 15 * mm

#         c.setLineWidth(1)
#         c.roundRect(x0, y - card_height, card_width, card_height, 5 * mm, stroke=1, fill=0)

#         header_h = 10 * mm
#         c.setFillColor(colors.lightgrey)
#         c.roundRect(x0, y - header_h, card_width, header_h, 5 * mm, stroke=0, fill=1)
#         c.setFillColor(colors.black)
#         c.setFont("Helvetica-Bold", 12)
#         c.drawString(x0 + 5 * mm, y - header_h + 2 * mm, f"KYC #{idx}")

#         labels = [
#             ("Name", kyc.name),
#             ("Father's Name", kyc.fathername or "—"),
#             ("Mobile", kyc.mobile_number),
#             ("Aadhar", kyc.aadhar_number),
#             ("Address", kyc.address),
#             ("Profession", kyc.profession or "—"),
#             ("Contact SH", kyc.contactSH or "—"),
#             ("Name SH", kyc.nameSH or "—"),
#             ("Investment", str(kyc.investmentamt) if kyc.investmentamt else "—"),
#         ]

#         col_x = [x0 + 5 * mm, x0 + card_width / 2 + 5 * mm]
#         c.setFont("Helvetica", 10)
#         line_h = 6 * mm
#         start_y = y - header_h - 5 * mm

#         for i, (label, val) in enumerate(labels):
#             col = i % 2
#             row = i // 2
#             text_y = start_y - row * line_h
#             c.drawString(col_x[col], text_y, f"{label}: {val}")

#         y -= card_height + 5 * mm

#     c.save()
#     return response




# from reportlab.lib.pagesizes import A4
# from reportlab.pdfgen import canvas
# from reportlab.lib import colors
# from reportlab.lib.units import mm
# from django.http import HttpResponse
# from django.contrib.auth.decorators import login_required
# from .models import MyKYC, SubKYC

# @login_required
# def download_kyc_pdf(request, kyc_type):
#     response = HttpResponse(content_type='application/pdf')
#     filename = "my_kyc_report.pdf" if kyc_type == 'my' else "sub_kyc_report.pdf"
#     response['Content-Disposition'] = f'attachment; filename="{filename}"'

#     c = canvas.Canvas(response, pagesize=A4)
#     W, H = A4
#     margin = 20 * mm
#     card_width = W - 2 * margin
#     card_height = 60 * mm
#     x0 = margin
#     y = H - margin

#     c.setFont("Helvetica-Bold", 18)
#     c.drawCentredString(W / 2, y, "KYC Report")
#     y -= 15 * mm

#     # Select model based on kyc_type
#     user = request.user
#     if kyc_type == 'my':
#         if user.is_superuser or getattr(user, "is_main_user", False):
#             kyc_list = MyKYC.objects.all()
#         elif getattr(user, "is_sub_mainuser", False):
#             kyc_list = MyKYC.objects.filter(user__parent=user)
#         else:
#             kyc_list = MyKYC.objects.filter(user=user)
#     else:
#         if user.is_superuser or getattr(user, "is_main_user", False):
#             kyc_list = SubKYC.objects.all()
#         elif getattr(user, "is_sub_mainuser", False):
#             kyc_list = SubKYC.objects.filter(user__parent=user)
#         else:
#             kyc_list = SubKYC.objects.filter(created_by=user).exclude(user=user)

#     for idx, kyc in enumerate(kyc_list, 1):
#         if y - card_height < margin:
#             c.showPage()
#             y = H - margin
#             c.setFont("Helvetica-Bold", 18)
#             c.drawCentredString(W / 2, y, "KYC Report")
#             y -= 15 * mm

#         c.setLineWidth(1)
#         c.roundRect(x0, y - card_height, card_width, card_height, 5 * mm, stroke=1, fill=0)

#         header_h = 10 * mm
#         c.setFillColor(colors.lightgrey)
#         c.roundRect(x0, y - header_h, card_width, header_h, 5 * mm, stroke=0, fill=1)
#         c.setFillColor(colors.black)
#         c.setFont("Helvetica-Bold", 12)
#         c.drawString(x0 + 5 * mm, y - header_h + 2 * mm, f"KYC #{idx}")

#         labels = [
#             ("Name", kyc.name),
#             ("Father's Name", kyc.fathername or "—"),
#             ("Mobile", kyc.mobile_number),
#             ("Aadhar", kyc.aadhar_number),
#             ("Address", kyc.address),
#             ("Profession", kyc.profession or "—"),
#             ("Contact SH", kyc.contactSH or "—"),
#             ("Name SH", kyc.nameSH or "—"),
#             ("Investment", str(kyc.investmentamt) if kyc.investmentamt else "—"),
#         ]

#         col_x = [x0 + 5 * mm, x0 + card_width / 2 + 5 * mm]
#         c.setFont("Helvetica", 10)
#         line_h = 6 * mm
#         start_y = y - header_h - 5 * mm

#         for i, (label, val) in enumerate(labels):
#             col = i % 2
#             row = i // 2
#             text_y = start_y - row * line_h
#             c.drawString(col_x[col], text_y, f"{label}: {val}")

#         y -= card_height + 5 * mm

#     c.save()
#     return response



################################################################
# subkyc excel and pdf

from django.shortcuts import render
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import MyKYC, SubKYC, BondImage
from django.contrib.auth import get_user_model

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.units import mm

import openpyxl
from openpyxl.styles import Font, Alignment

User = get_user_model()


@login_required
def download_subkyc_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="sub_kyc_report.pdf"'

    c = canvas.Canvas(response, pagesize=A4)
    W, H = A4
    margin = 20 * mm
    card_width = W - 2 * margin
    card_height = 60 * mm
    x0 = margin
    y = H - margin

    c.setFont("Helvetica-Bold", 18)
    c.drawCentredString(W / 2, y, "Sub-KYC Report")
    y -= 15 * mm

    user = request.user
    hidden_sub_kyc_ids = request.session.get('hidden_sub_kyc', [])

    if user.is_superuser or user.is_main_user:
        kyc_list = SubKYC.objects.filter(created_by=user) | SubKYC.objects.filter(user=user)
    else:
        kyc_list = SubKYC.objects.filter(user=user)

    kyc_list = kyc_list.exclude(id__in=hidden_sub_kyc_ids).distinct()

    for idx, kyc in enumerate(kyc_list, 1):
        if y - card_height < margin:
            c.showPage()
            y = H - margin
            c.setFont("Helvetica-Bold", 18)
            c.drawCentredString(W / 2, y, "Sub-KYC Report")
            y -= 15 * mm

        c.setLineWidth(1)
        c.roundRect(x0, y - card_height, card_width, card_height, 5 * mm, stroke=1, fill=0)

        header_h = 10 * mm
        c.setFillColor(colors.lightgrey)
        c.roundRect(x0, y - header_h, card_width, header_h, 5 * mm, stroke=0, fill=1)
        c.setFillColor(colors.black)
        c.setFont("Helvetica-Bold", 12)
        c.drawString(x0 + 5 * mm, y - header_h + 2 * mm, f"Sub-KYC #{idx}")

        labels = [
            ("Name", kyc.name),
            ("Father's Name", kyc.fathername or "—"),
            ("Mobile", kyc.mobile_number),
            ("Aadhar", kyc.aadhar_number),
            ("Address", kyc.address),
            ("Profession", kyc.profession or "—"),
            ("Contact SH", kyc.contactSH or "—"),
            ("Name SH", kyc.nameSH or "—"),
            ("Investment", str(kyc.investmentamt) if kyc.investmentamt else "—"),
        ]

        col_x = [x0 + 5 * mm, x0 + card_width / 2 + 5 * mm]
        c.setFont("Helvetica", 10)
        line_h = 6 * mm
        start_y = y - header_h - 5 * mm

        for i, (label, val) in enumerate(labels):
            col = i % 2
            row = i // 2
            text_y = start_y - row * line_h
            c.drawString(col_x[col], text_y, f"{label}: {val}")

        y -= card_height + 5 * mm

    c.save()
    return response



@login_required
def download_subkyc_excel(request):
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="sub_kyc_report.xlsx"'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sub-KYC Report"

    headers = [
        "Name", "Father's Name", "Mobile", "Aadhar",
        "Address", "Profession", "Contact SH",
        "Name SH", "Investment"
    ]
    ws.append(headers)

    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    user = request.user
    hidden_sub_kyc_ids = request.session.get('hidden_sub_kyc', [])

    if user.is_superuser or user.is_main_user:
        kyc_list = SubKYC.objects.filter(created_by=user) | SubKYC.objects.filter(user=user)
    else:
        kyc_list = SubKYC.objects.filter(user=user)

    kyc_list = kyc_list.exclude(id__in=hidden_sub_kyc_ids).distinct()

    for kyc in kyc_list:
        ws.append([
            kyc.name,
            kyc.fathername or "",
            kyc.mobile_number,
            kyc.aadhar_number,
            kyc.address,
            kyc.profession or "",
            kyc.contactSH or "",
            kyc.nameSH or "",
            kyc.investmentamt if kyc.investmentamt else 0,
        ])

    wb.save(response)
    return response



@login_required
def profile_view(request):
    user = request.user

    if request.method == 'POST':
        user.username = request.POST['username']
        user.email = request.POST['email']
        user.phone_number = request.POST['phone_number']
        user.save()
        messages.success(request, "Profile updated successfully.")
        return redirect('profile')

    return render(request, 'profile.html', {'user_obj': user})



from django.contrib.auth.decorators import login_required
from django.shortcuts import redirect
from django.contrib import messages
from django.utils import timezone
import datetime

from .models import MyKYC, BondImage, User


@login_required
def add_my_kyc(request):
    if request.method == "POST":
        created_by = request.user

        # Validate selected user
        try:
            data_for_user = User.objects.get(id=request.POST.get('data_for_user'))
        except User.DoesNotExist:
            messages.error(request, "Invalid user selected.")
            return redirect("formpage")

        # Convert numeric fields safely
        def to_int(val):
            try:
                return int(val) if val else None
            except ValueError:
                return None

        my_kyc = MyKYC.objects.create(
            user=data_for_user,
            created_by=created_by,
            membershipno=to_int(request.POST.get('membershipno')),
            # membershiptype=request.POST.get('membershiptype'),
            depositorsname=request.POST.get('depositorsname'),
            depositorsaddress=request.POST.get('depositorsaddress'),
            bondholdername=request.POST.get('bondholdername'),
            # projectname=request.POST.get('projectname'),
            depositormobile_number=request.POST.get('depositormobile_number'),
            agentname=request.POST.get('agentname'),
            agentaddress=request.POST.get('agentaddress'),
            nameofdirector=request.POST.get('nameofdirector'),
            aadhar_number=request.POST.get('aadhar_number'),
            pan_number=request.POST.get('pan_number'),
            ration_number=request.POST.get('ration_number'),
            bankname=request.POST.get('bankname'),
            bankaccno=request.POST.get('bankaccno'),
            ifscno=request.POST.get('ifscno'),
            aadhar_front_image=request.FILES.get('aadhar_front_image'),
            aadhar_back_image=request.FILES.get('aadhar_back_image'),
            passportphoto=request.FILES.get('passportphoto'),
        )

        # Save bond images and related data
        bond_files = request.FILES.getlist('bonds')
        for i, bond_file in enumerate(bond_files):
            BondImage.objects.create(
                my_kyc=my_kyc,
                bondholdername=request.POST.get(f'bondholdername_{i}'),
                image=bond_file,
                companyname=request.POST.get(f'companyname_{i}'),
                projectname=request.POST.get(f'projectname_{i}'),
                refundamount=request.POST.get(f'refundamount_{i}'),
                balanceamount=request.POST.get(f'balanceamount_{i}'),
                amount=to_int(request.POST.get(f'amount_{i}')) or 0,
                investment_date=request.POST.get(f'investment_date_{i}') or datetime.date.today(),
                customer_id=request.POST.get(f'customer_id_{i}')
            )

        messages.success(request, "My KYC data saved successfully.")
        return redirect('formpage')

from django.contrib.auth.decorators import login_required
from django.shortcuts import redirect
from django.contrib import messages
from .models import SubKYC, BondImage, User

@login_required
def add_other_kyc(request):
    if request.method == "POST":
        created_by = request.user

        try:
            data_for_user = User.objects.get(id=request.POST.get('data_for_user'))
        except User.DoesNotExist:
            messages.error(request, "Invalid user selected.")
            return redirect("formpage")

        try:
            investmentamt = int(request.POST.get('investmentamt')) if request.POST.get('investmentamt') else None
        except ValueError:
            messages.error(request, "Investment amount must be numeric.")
            return redirect("formpage")

        kyc = SubKYC.objects.create(
            user=data_for_user,
            created_by=created_by,
            name=request.POST.get('name'),
            age=request.POST.get('age'),
            mobile_number=request.POST.get('mobile_number'),
            fathername=request.POST.get('fathername'),
            address=request.POST.get('address'),
            aadhar_number=request.POST.get('aadhar_number'),
            profession=request.POST.get('profession'),
            contactSH=request.POST.get('contactSH'),
            nameSH=request.POST.get('nameSH'),
            investmentamt=investmentamt,
            aadhar_front_image=request.FILES.get('aadhar_front_image'),
            aadhar_back_image=request.FILES.get('aadhar_back_image'),
            passportphoto=request.FILES.get('passportphoto'),
        )

        for bond_file in request.FILES.getlist('bonds'):
            BondImage.objects.create(sub_kyc=kyc, image=bond_file)  # ✅ FIXED HERE

        messages.success(request, "Sub-KYC added successfully.")
        return redirect('formpage')

